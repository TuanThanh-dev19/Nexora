from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

WORK = Path(__file__).resolve().parent
OUT = WORK.parent / 'ho-so-nop-nexora-20261029'
OUT.mkdir(parents=True, exist_ok=True)
facts = json.loads((WORK / 'sources.json').read_text(encoding='utf-8'))
repo = facts['repository']
sheet = facts['backlog']

txt = [
    'THÔNG TIN ĐỀ TÀI VÀ THÀNH VIÊN', '',
    'Tên topic: ' + facts['topic'],
    'Môn học: ' + facts['course'], '',
    'Thành viên:',
]
for idx, (name, account) in enumerate(facts['members'], 1):
    txt.append(f'{idx}. {name} | GitHub: {account}')
txt += ['', 'Link Git project: ' + repo, '', 'Ngày cập nhật: 29/10/2026', '']
(OUT / '01_Thong_tin_de_tai_va_thanh_vien.txt').write_text('\n'.join(txt), encoding='utf-8-sig', newline='\r\n')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.68)
sec.left_margin = sec.right_margin = Inches(0.75)
sec.footer_distance = Inches(0.3)

for sty in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet']:
    st = doc.styles[sty]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    fonts = st.element.get_or_add_rPr().rFonts
    fonts.set(qn('w:cs'), 'Times New Roman')
    for attribute in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme']:
        fonts.attrib.pop(qn('w:' + attribute), None)
    for border in st.element.findall('.//' + qn('w:pBdr')):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(7)
doc.styles['Subtitle'].font.size = Pt(13)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(12)
for name, size in [('Heading 1',15), ('Heading 2',12)]:
    st=doc.styles[name]
    st.font.size=Pt(size)
    st.font.bold=True
    st.paragraph_format.space_before=Pt(10)
    st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.keep_with_next=True

def para(text='', bold=False, style=None):
    p=doc.add_paragraph(style=style)
    p.add_run(text).bold=bold
    return p

def heading(text):
    return doc.add_paragraph(text, style='Heading 1')

def link(p, text, url):
    rel=p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'),'Times New Roman'); fonts.set(qn('w:hAnsi'),'Times New Roman'); pr.append(fonts)
    color=OxmlElement('w:color'); color.set(qn('w:val'),'1F4E79'); pr.append(color)
    under=OxmlElement('w:u'); under.set(qn('w:val'),'single'); pr.append(under)
    r.append(pr); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)

def table(headers, rows, widths, font=11):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c, width in zip(t.columns,widths): c.width=Inches(width)
    for i,h in enumerate(headers): t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i, text in enumerate(row): cells[i].text=str(text)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),'80' if side in ['top','bottom'] else '90'); e.set(qn('w:type'),'dxa'); margins.append(e)
    t._tbl.tblPr.append(margins)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7EDF3'); c._tc.get_or_add_tcPr().append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1.05
                if ri==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(font);r.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
rr=footer.add_run('Nexora | ');rr.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

# Trang 1
doc.add_paragraph('Thông tin đề tài Nexora', style='Title')
doc.add_paragraph('Hệ thống báo cáo sự cố và điều phối bảo trì trong khuôn viên trường', style='Subtitle')
para('Môn học: SBA301 – Integrate Single Page Application with Spring Boot')
para('Báo cáo tiến độ đến ngày 29/10/2026')
para('Nhóm xây dựng Nexora để tiếp nhận sự cố về cơ sở vật chất, giao việc cho kỹ thuật viên và theo dõi kết quả xử lý. Báo cáo trình bày mục tiêu đề tài, phạm vi chức năng, phân công và tiến độ hiện tại để giảng viên theo dõi quá trình thực hiện.')
heading('1 Bối cảnh và mục tiêu')
para('Khi thông tin sự cố được trao đổi qua nhiều kênh, người báo khó theo dõi tiến độ, còn bộ phận vận hành khó xác định người phụ trách và lịch sử xử lý. Nexora tập trung các thông tin này vào một hệ thống, gắn mỗi Ticket với địa điểm, người báo, người xử lý và kết quả công việc.')
para('Nhóm hướng tới việc tạo phiếu bằng biểu mẫu phù hợp từng loại sự cố, phân công theo điều kiện kỹ thuật viên và lưu được nhật ký cùng ảnh kết quả để người báo kiểm tra trước khi đóng phiếu.')
heading('2 Người sử dụng và luồng nghiệp vụ')
table(['Vai trò','Nhu cầu chính'],[
('Requester','Tạo Ticket, xem tiến độ và ảnh kết quả, xác nhận đóng Ticket của mình.'),
('Manager','Tiếp nhận Ticket, kiểm tra kỹ thuật viên phù hợp, phân công và theo dõi xử lý.'),
('Technician','Xem việc được giao, bắt đầu xử lý, ghi nhật ký, tải ảnh kết quả và hoàn tất công việc.'),
('Admin','Quản lý tài khoản, vai trò, danh mục và biểu mẫu cấu hình hệ thống.')
],[1.1,5.9])
para('Luồng chính: Tạo Ticket → Tiếp nhận → Phân công → Bắt đầu xử lý → Ghi nhật ký và ảnh kết quả → Hoàn tất xử lý → Người báo xác nhận đóng Ticket.')
heading('3 Công nghệ và tổ chức dữ liệu')
para('Frontend sử dụng React và Vite; backend sử dụng Java 21 và Spring Boot 3. SQL Server quản lý dữ liệu nghiệp vụ có quan hệ; MongoDB lưu các document phụ trợ. Flyway quản lý phiên bản cơ sở dữ liệu, Docker Compose hỗ trợ dựng môi trường và GitHub quản lý mã nguồn, issue cùng pull request.')

# Trang 2
doc.add_page_break()
heading('4 Thành viên và phân công chính')
para('Phân công dưới đây phản ánh nhóm nghiệp vụ và các issue được giao trong backlog. Những phần có nhiều người tham gia được theo dõi chi tiết bằng issue GitHub, thay vì quy toàn bộ một module cho một người.')
table(['Thành viên','Tài khoản GitHub','Phần việc chính'],[
('Trần Phan Nhật Huy','Huytran2903','Đăng nhập và phiên làm việc; quản lý user, role và cấu hình Technician; kiểm tra eligibility; kiểm thử hồi quy API.'),
('Lâm Thoại Thanh','thoaithanh-web','Category và dynamic form; tạo Ticket và validation; tìm kiếm Ticket theo mã hoặc tiêu đề.'),
('Phan Tuấn Thành','TuanThanh-dev19','Ticket read và workflow; review, assign, close; upload ảnh kết quả; reassign, kế hoạch unassign và kiểm thử E2E.'),
('Tô Phi Oanh','OanhPhiTo','Location, equipment type và asset; My Work, start, work log và resolve; thông báo trong ứng dụng.'),
('Trần Quốc Thái','tranthaigin','Hạ tầng file; xem và tải ảnh; Dashboard, SLA và báo cáo Ticket xuất CSV; hỗ trợ tích hợp.')
],[1.48,1.25,4.27])
heading('5 Phạm vi đã triển khai')
for text in [
    'Đăng nhập, quản lý phiên làm việc và phân quyền; quản trị tài khoản, role, Technician, địa điểm, loại thiết bị và asset.',
    'Quản lý Category và biểu mẫu có phiên bản; tạo Ticket theo published form, kiểm tra dữ liệu động và thông tin thiết bị hợp lệ.',
    'Đọc danh sách và chi tiết Ticket; review và phân công; My Work, bắt đầu xử lý, nhật ký, upload ảnh kết quả, resolve và reporter close.',
    'Xem và tải ảnh có phân quyền; Dashboard theo phạm vi người dùng, thông tin SLA; tìm kiếm Ticket theo mã và tiêu đề tại issue #95.'
]: para(text, style='List Bullet')
para('Các hạng mục này đã có issue triển khai đóng trên GitHub. Bộ kiểm thử hồi quy API và E2E còn đang mở, vì vậy việc đóng issue triển khai chưa được dùng để khẳng định toàn bộ luồng đã nghiệm thu.')
heading('6 Kiểm soát nghiệp vụ')
para('Hệ thống kiểm tra quyền theo vai trò và quan hệ với từng Ticket. Các thao tác xử lý yêu cầu đúng người được phân công và đúng trạng thái; thay đổi workflow có kiểm soát version. Ticket giữ phiên bản form khi tạo để đọc lại lịch sử. Ảnh kết quả là tùy chọn và được truy cập qua API có kiểm tra quyền.')

# Trang 3
doc.add_page_break()
heading('7 Tiến độ và phần việc tiếp theo')
para('Tại ngày 29/10/2026, backlog có 57 dòng công việc: 51 Done, 3 In Progress và 3 To Do. Bảng Usecase gồm 44 usecase trong phạm vi đề tài. Đây là số dòng theo dõi, không phải số usecase đã nghiệm thu hoặc tỷ lệ công sức của thành viên.')
table(['Nguồn','Phần việc còn lại','Phụ trách','Hạn dự kiến','Trạng thái'],[
('#90','Kiểm thử hồi quy API Ticket','Trần Phan Nhật Huy','31/10/2026','To Do'),
('#91','Kiểm thử E2E Ticket với API thật','Phan Tuấn Thành','05/11/2026','To Do'),
('#92','Reassign Technician','Phan Tuấn Thành','03/11/2026','In Progress'),
('#93','Báo cáo Ticket và xuất CSV','Trần Quốc Thái','30/10/2026','In Progress'),
('#94','Thông báo trong ứng dụng','Tô Phi Oanh','30/10/2026','In Progress'),
('Kế hoạch','Unassign Technician','Phan Tuấn Thành','04/11/2026','To Do')
],[0.65,2.4,1.5,1.15,1.3],font=10.5)
para('Unassign Technician đã được đưa vào phạm vi thực hiện; người phụ trách và hạn 04/11 là kế hoạch, chưa có issue triển khai riêng. Các hạn trong bảng theo lịch hiện tại của Product Backlog; milestone Sprint 4 trên GitHub vẫn kết thúc ngày 30/10 và cần được thống nhất với lịch điều chỉnh.')
para('Nhóm ưu tiên chạy hồi quy API và E2E của luồng MVP hiện có, sau đó bổ sung ca bị ảnh hưởng khi reassign và unassign được tích hợp. Dự kiến hoàn thiện unassign ngày 04/11 và dành ngày 05/11 để hoàn tất kiểm thử, xử lý lỗi và rà soát tài liệu.')
heading('8 Phạm vi chưa thực hiện')
para('Reject Ticket, Cancel Ticket và Reopen Ticket đã được loại khỏi danh mục thực hiện. Comments và feedback chưa thuộc luồng demo hoàn chỉnh. Nhóm tiếp tục theo dõi tiến độ các chức năng báo cáo, thông báo và điều phối bổ sung bằng các issue đang mở.')
heading('9 Mã nguồn và tài liệu theo dõi')
p=para();p.add_run('Git project: ');link(p,repo,repo)
p=para();p.add_run('Backlog và danh mục usecase: ');link(p,'Nexora Product Backlog',sheet)
p=para();p.add_run('Issue đang mở: ')
for k,n in enumerate([90,91,92,93,94]):
    if k:p.add_run('  |  ')
    link(p,f'#{n}',f'{repo}/issues/{n}')
para('Tài liệu trong repository: README.md, docs/PROJECT_REPORT.md, docs/ARCHITECTURE.md, docs/API_CONVENTIONS.md và docs/TESTING.md.')

doc.core_properties.title='Thông tin đề tài và tiến độ Nexora'
doc.core_properties.subject='Báo cáo đề tài SBA301 ngày 29/10/2026'
doc.core_properties.author='Nhóm dự án Nexora'
doc.core_properties.keywords='Nexora, SBA301, báo cáo tiến độ'
doc.core_properties.comments=''
for paragraph in doc.paragraphs:
    for border in paragraph._p.findall('.//' + qn('w:pBdr')):
        border.getparent().remove(border)
    for run in paragraph.runs:
        run.font.name='Times New Roman'
        run._r.get_or_add_rPr().rFonts.set(qn('w:cs'),'Times New Roman')
path=OUT/'02_Thong_tin_de_tai_va_tien_do.docx'
doc.save(path)
print(json.dumps({'txt':str(OUT/'01_Thong_tin_de_tai_va_thanh_vien.txt'),'docx':str(path),'members':len(facts['members'])},ensure_ascii=False))
