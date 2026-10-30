import json
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

root = Path(__file__).parent
plan = json.loads((root / 'update-payload.json').read_text(encoding='utf-8'))
before = openpyxl.load_workbook(root / 'before.xlsx')
after = openpyxl.load_workbook(root / 'after.xlsx')
cached = openpyxl.load_workbook(root / 'after.xlsx', data_only=True)

allowed = {
    'Product Backlog': {(r, c) for r in range(52, 58) for c in range(1, 10)} | {(46, 9), (49, 9)},
    'Config': {(r, c) for r in range(36, 41) for c in range(1, 4)} | {(2, 2), (2, 3), (3, 11), (7, 11), (10, 11)},
    'Import Audit': {(r, c) for r in range(90, 96) for c in range(1, 14)},
}
changes = Counter()
for sheet_name in before.sheetnames:
    old, new = before[sheet_name], after[sheet_name]
    for row in range(1, max(old.max_row, new.max_row) + 1):
        for col in range(1, max(old.max_column, new.max_column) + 1):
            if old.cell(row, col).value != new.cell(row, col).value:
                assert (row, col) in allowed[sheet_name], (sheet_name, row, col, 'unexpected value change')
                changes[sheet_name] += 1

sheet = after['Product Backlog']
assert sheet.tables['ProductBacklogTable'].ref == 'A1:I57'
for row, expected in enumerate(plan['backlogRows'], 52):
    for col, value in enumerate(expected, 1):
        if col == 6:
            value = datetime.strptime(value, '%d/%m/%Y')
        assert sheet.cell(row, col).value == value, (row, col, sheet.cell(row, col).value, value)
    assert sheet.cell(row, 6).number_format == 'dd/MM/yyyy'
assert 'D2:D57' in {str(v.sqref) for v in sheet.data_validations.dataValidation}
assert 'E2:E57' in {str(v.sqref) for v in sheet.data_validations.dataValidation}
assert after['Config'].tables['ConfigUsecase'].ref == 'A1:C40'
assert after['Import Audit'].tables['ImportAuditTable'].ref == 'A1:M95'
for row, expected in enumerate(plan['configRows'], 36):
    for col, value in enumerate(expected, 1):
        if col == 3:
            value = datetime.strptime(value, '%d/%m/%Y')
        assert after['Config'].cell(row, col).value == value
assert after['Config']['B2'].value == 'Lâm Thoại Thanh'
assert after['Config']['C2'].value == datetime(2026, 10, 30)
for cell, value in plan['configGuides'].items():
    assert after['Config'][cell].value == value
for cell, value in plan['backlogNoteUpdates'].items():
    assert sheet[cell].value == value
for row, expected in enumerate(plan['auditRows'], 90):
    for col, value in enumerate(expected, 1):
        if col == 11:
            value = datetime.strptime(value, '%d/%m/%Y')
        assert (after['Import Audit'].cell(row, col).value or '') == value

counts = Counter(sheet.cell(r, 5).value for r in range(2, 58))
assert sum(counts.values()) == 56
for row in range(2, 7):
    member = sheet.cell(row, 11).value
    assert cached['Product Backlog'].cell(row, 12).value == counts[member]
    assert abs(cached['Product Backlog'].cell(row, 13).value - counts[member] / 56) < 1e-9
assert cached['Product Backlog']['L7'].value == 56
assert abs(cached['Product Backlog']['M7'].value - 1) < 1e-9

statuses = Counter(sheet.cell(r, 4).value for r in range(2, 58))
assert statuses == {'Done': 50, 'In Progress': 1, 'To Do': 5}, statuses
new_names = {after['Config'].cell(r, 1).value for r in range(2, 41)}
for row in range(52, 58):
    assert set(sheet.cell(row, 2).value.split('; ')) <= new_names

print(json.dumps({'verified': True, 'value_changes': changes, 'backlog_rows': 56,
                  'statuses': statuses, 'member_counts': counts,
                  'tables_and_dropdowns_extended': True, 'unrelated_cells_preserved': True},
                 ensure_ascii=False, indent=2))
